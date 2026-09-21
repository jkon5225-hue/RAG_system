# streamlit run login.py
# データベース丸ごと消去　rmdir /s /q chroma_db2
# ubuntuでWORDをPDFに変換するには事前にコマンド:  sudo apt install libreoffice 
# windowsなら　pip install docx2pdf

import streamlit as st
from auth import require_admin
from auth import sidebar_user
from pptx import Presentation

require_admin()

sidebar_user()

import os
import re
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"

import fitz
import chromadb
import uuid
from sentence_transformers import SentenceTransformer
import ollama
from st_copy_to_clipboard import st_copy_to_clipboard
import streamlit.components.v1 as components
import tempfile
import subprocess
from io import BytesIO
import platform
from docx2pdf import convert
import unicodedata

from db import get_question_conn
import time
from datetime import datetime
import requests
from urllib.parse import quote

# =====================================================
# Settings
# =====================================================

BASE_DB = "./chroma_db2"
EMBED_MODEL = "intfloat/multilingual-e5-large"
LLM_MODEL = "granite4.1:8b"

BASE_PDF = "./pages/lite3"

GENRE_SETTINGS = {
    "議事録":{
        "db":"minutes","pdf":"minutes",
        "chunk_size":400,"overlap":80,"top_k":5,
        "prompt": """
        あなたは議事録検索AIです。提出された議事録の文脈のみを根拠に、事実に基づき日本語で回答してください。
    
        【回答ルール】
        ・「却下」「承認」「保留」「決定」などのステータスは、議事録の記載通りに正確に伝えてください。
        ・推測や、文脈にない情報を付け加えて回答することを厳禁とします。
        ・該当する情報が文脈内に見当たらない場合は、「議事録に記載がありません」と回答してください。
        """,
        "file_types":["pdf", "docx"]
        },
    "規約":{
        "db":"rules","pdf":"rules",
        "chunk_size":400,"overlap":80,"top_k":4,
        "prompt":"""
        あなたは規約検索AIです。提出された条文のみを根拠に日本語で回答してください。
        
        【回答ルール】
        ・文書に記載されている内容のみを回答する。
        ・推測や一般論を付け加えない。
        ・手続き、申請先、連絡先、対象者、期限、必要書類などは文書の記載どおりに回答する。
        ・複数の文書に情報がある場合は整理してまとめる。
        """,
        "file_types":["pdf"]
        },
    "PowerPoint":{
        "db":"powerpoint",
        "pdf":"powerpoint",
        "chunk_size":0,      # PowerPointでは使用しない
        "overlap":0,         # PowerPointでは使用しない
        "top_k":5,
        "prompt":"""
        あなたはPowerPoint資料検索AIです。

        質問に対して、関連性の高いPowerPointスライドを優先して回答してください。
        回答には資料名とスライド番号を示し、推測で内容を補わないでください。
        """,
        "file_types":["pptx"]
        },
    "生活関連":{
        "db":"life","pdf":"life",
        "chunk_size":400,"overlap":80,"top_k":5,
        "prompt":"""
        あなたは社内生活便利帳検索AIです。

        提出された文書のみを根拠に日本語で回答してください。

        【回答ルール】
        ・文書に記載されている内容のみを回答する。
        ・推測や一般論を付け加えない。
        ・手続き、申請先、連絡先、対象者、期限、必要書類などは文書の記載どおりに回答する。
        ・複数の文書に情報がある場合は整理してまとめる。
        ・文書に記載がない場合は「資料には記載がありません」と回答する。
        """,
        "file_types":["pdf"]
        },

    "論文":{
        "db":"papers","pdf":"papers",
        "chunk_size":1000,"overlap":200,"top_k":5,
        "prompt":"""
        あなたは論文翻訳・要約の専門家です。
        【用語ルール】
        ・専門用語は一般的な日本語訳が存在する場合のみ翻訳する。
        ・一般的でない用語は英語表記を維持する。
        ・意味のないカタカナ音訳は禁止する。
        ・初出では「日本語（英語）」で表記する。
        ・略語は初出で正式名称を書く。
        ・タンパク質名、遺伝子名、ペプチド名、材料名、試薬名は原則英語表記を維持する。

        【回答ルール】
        ・論文の内容を忠実に日本語で説明する。
        ・推測で訳語を作らない。
        ・原文にない情報は付け加えない。
        """,
        "file_types":["pdf"]
        }
}

# =====================================================
# Models
# =====================================================

@st.cache_resource
def load_embedder():
    return SentenceTransformer(EMBED_MODEL)

embedder = load_embedder()
# 登録エリア＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
def chunk_by_section(text):
    # 「3.」「4.」「5.」などで区切る
    parts = re.split(r'(?=^\d+\.\s)', text, flags=re.MULTILINE)

    return [p.strip() for p in parts if p.strip()]
# =====================================================
# PDF
# =====================================================
def pdf_to_text(pdf_path):

    pdf = fitz.open(pdf_path)

    pages = []

    for page_no, page in enumerate(pdf):

        # 通常テキスト
        text = page.get_text()

        # ==========================
        # 表をMarkdown化して追加
        # ==========================
        try:

            tables = page.find_tables()

            if tables.tables:

                text += "\n\n"

                for table in tables:

                    data = table.extract()

                    if not data:
                        continue

                    # ---------- ヘッダ ----------
                    header = [
                        "" if v is None else str(v).replace("\n", " ")
                        for v in data[0]
                    ]

                    text += "|" + "|".join(header) + "|\n"
                    text += "|" + "|".join(["---"] * len(header)) + "|\n"

                    # ---------- データ ----------
                    for row in data[1:]:

                        row = [
                            "" if v is None else str(v).replace("\n", " ")
                            for v in row
                        ]

                        text += "|" + "|".join(row) + "|\n"

                    text += "\n"

        except Exception:
            # 表が無い・古いPyMuPDFでもそのまま続行
            pass

        pages.append((page_no, text))

    pdf.close()

    return pages

def chunk_by_length(text,size,overlap):
    out=[]
    p=0
    while p<len(text):
        c=text[p:p+size]
        if len(c.strip())>50:
            out.append(c)
        p+=size-overlap
    return out


def chunk_text(pages, genre):

    s = GENRE_SETTINGS[genre]
    chunks = []

    # -----------------------------
    # 規約
    # -----------------------------
    if genre == "規約":

        all_text = ""
        page_map = []

        for page_id, text in pages:
            page_map.append((len(all_text), page_id))
            all_text += "\n" + text

        for c in chunk_by_section(all_text):

            pos = all_text.find(c)

            page = 0
            for start, page_id in page_map:
                if start <= pos:
                    page = page_id
                else:
                    break

            chunks.append({
                "text": c,
                "page": page,
                "id": str(uuid.uuid4())
            })

        return chunks

    # -----------------------------
    # PowerPoint
    # 1スライド＝1チャンク
    # -----------------------------
    elif genre == "PowerPoint":

        for slide_id, text in pages:

            if len(text.strip()) == 0:
                continue

            chunks.append({
                "text": text,
                "page": slide_id,     # pageをそのままスライド番号として使用
                "id": str(uuid.uuid4())
            })

        return chunks

    # -----------------------------
    # PDF（従来）
    # -----------------------------
    else:

        for page_id, text in pages:

            parts = chunk_by_length(
                text,
                s["chunk_size"],
                s["overlap"]
            )

            for c in parts:
                chunks.append({
                    "text": c,
                    "page": page_id,
                    "id": str(uuid.uuid4())
                })

        return chunks
# =====================================================
# PowerPoint to TEXT
# =====================================================
def ppt_to_text(uploaded_file):

    uploaded_file.seek(0)

    prs = Presentation(uploaded_file)

    slides = []

    for slide_no, slide in enumerate(prs.slides):

        texts = []

        for shape in slide.shapes:

            # テキストを持つShapeのみ取得
            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    texts.append(text)

        # 1スライド＝1チャンク
        slide_text = "\n".join(texts)

        slides.append(
            (
                slide_no,      # 0始まり（表示時に+1）
                slide_text
            )
        )

    return slides
# =========================
# OCR判定
# =========================
def ocr_pdf(uploaded_file):
    uploaded_file.seek(0)

    pdf = fitz.open(
        stream=uploaded_file.read(),
        filetype="pdf"
    )

    OCR_THRESHOLD = 100

    dst = fitz.open()

    for i, page in enumerate(pdf):
        text = page.get_text().strip()
        # 文字情報が少ないページだけOCR
        if len(text) < OCR_THRESHOLD:
            # ==================
            # TXT数が少なかったらOCR化
            # ==================
            # 画像化
            pix = page.get_pixmap(dpi=400)

            # OCR付きPDF生成
            pdfbytes = pix.pdfocr_tobytes(
                language="jpn+eng"
            )

            # 一時PDFを開く
            ocr_doc = fitz.open("pdf", pdfbytes)

            # OCRテキスト取得
            ocr_text = ocr_doc[0].get_text()

            # 正規化
            ocr_text = unicodedata.normalize("NFKC", ocr_text)

            # 出力へ追加（Documentを渡す）
            dst.insert_pdf(
                ocr_doc,
                from_page=0,
                to_page=0
            )
            ocr_doc.close()
        else:
            # そのまま出力へ追加
            dst.insert_pdf(
                pdf,
                from_page=i,
                to_page=i
            )
    
    # =========================
    # PDF保存
    # =========================
    pdf.close()
    pdf_path = os.path.join(PDF_DIR, uploaded_file.name)
    dst.save(
        pdf_path,
        garbage=4,
        deflate=True
    )
    dst.close()

    return
# =====================================================
# Register
# =====================================================
def add_pdf(uploaded_file):

    # =========================
    # 既存PDFチェック
    # =========================
    existing = collection.get(
        where={"file": uploaded_file.name}
    )

    if existing and len(existing.get("ids", [])) > 0:
        print(f"Skip: {uploaded_file.name} already exists")
        return 0
    
    #=========================
    #画像PDFページOCR化
    #=========================
    ocr_pdf(uploaded_file)

    # =========================
    # PDF解析
    # =========================
    pdf_path = os.path.join(PDF_DIR, uploaded_file.name)
    pages = pdf_to_text(pdf_path)
    chunks = chunk_text(pages, genre)
    print("chunks =", len(chunks))

    texts = [c["text"] for c in chunks]
    if len(texts) == 0:
        st.error("chunkが0件です（PDF抽出失敗）")
        return 0

    embeddings = embedder.encode(
        texts,
        batch_size=64,   # または128
        normalize_embeddings=True,
        show_progress_bar=True
    )

    ids = [str(uuid.uuid4()) for _ in chunks]

    metadatas = [
        {
            "file": uploaded_file.name,
            "chunk": i,
            "page": chunks[i]["page"]
        }
        for i in range(len(chunks))
    ]

    # =========================
    # Chroma登録
    # =========================
    collection.add(
        ids=ids,
        documents=texts,   # ★ここ重要（chunksではない）
        embeddings=embeddings.tolist(),
        metadatas=metadatas
    )

    return len(chunks)

#PowerPoint登録 
def add_ppt(uploaded_file):
    # =========================
    # 既存PowerPointチェック
    # =========================
    existing = collection.get(
        where={"file": uploaded_file.name}
    )

    if existing and len(existing.get("ids", [])) > 0:
        print(f"Skip: {uploaded_file.name} already exists")
        return 0
    
    # =========================
    # PowerPoint保存
    # =========================
    ppt_path = os.path.join(PDF_DIR, uploaded_file.name)

    with open(ppt_path, "wb") as f:
        f.write(uploaded_file.getbuffer())

    # =========================
    # PowerPoint解析
    # =========================
    uploaded_file.seek(0)

    slides = ppt_to_text(uploaded_file)
    chunks = chunk_text(slides, genre)

    texts = [c["text"] for c in chunks]

    if len(texts) == 0:
        st.error("chunkが0件です（PowerPoint抽出失敗）")
        return 0

    embeddings = embedder.encode(
        texts,
        batch_size=64,   # または128
        normalize_embeddings=True,
        show_progress_bar=True
    )

    ids = [str(uuid.uuid4()) for _ in chunks]

    metadatas = [
        {
            "file": uploaded_file.name,
            "chunk": i,
            "page": chunks[i]["page"]      # ← スライド番号をpageに格納
        }
        for i in range(len(chunks))
    ]

    # =========================
    # Chroma登録
    # =========================
    collection.add(
        ids=ids,
        documents=texts,
        embeddings=embeddings.tolist(),
        metadatas=metadatas
    )

    return len(chunks)

# 登録エリア終わり＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊

# =====================================================
# Retrieve
# =====================================================
def retrieve(query, exclude_files=None):
    if exclude_files is None:
        exclude_files = []

    q_emb = embedder.encode(
        query,
        normalize_embeddings=True
    )

    results = collection.query(
        query_embeddings=[q_emb.tolist()],
        n_results=setting["top_k"]
    )
    docs = results["documents"][0]
    metas = results["metadatas"][0]
    distances = results["distances"][0]

    # ----------------------------
    # 除外
    # ----------------------------
    if exclude_files:

        new_docs = []
        new_metas = []
        new_distances = []

        for doc, meta, dist in zip(
            docs,
            metas,
            distances
        ):

            # ファイル名部分一致
            if any(
                key.lower() in meta["file"].lower()
                for key in exclude_files
            ):
                continue

            new_docs.append(doc)
            new_metas.append(meta)
            new_distances.append(dist)

        docs = new_docs
        metas = new_metas
        distances = new_distances

    return docs, metas, distances


@st.dialog("登録完了")
def show_upload_dialog(message):
    st.write(message)

    if st.button("OK"):
        st.session_state.uploader_key += 1
        st.rerun()

# =====================================================
# RAG
# =====================================================
def answer_question(question,exclude_files=None):

    docs, metas, distances = retrieve(
        question,
        exclude_files
    )

    context = "\n\n".join(docs)
    context = context[:2000]

    response = ollama.chat(
        model=LLM_MODEL,
        messages=[
            {
                "role":"system",
                "content":SYSTEM_PROMPT
            },
            {
                "role":"user",
                "content":f"""
                以下の文脈のみを使用してください。

                ====================

                {context}

                ====================

                質問
                {question}
                """
            }
        ], 
    )

    answer = response["message"]["content"]

    return answer, metas, docs, distances


def open_file(path):

    if os.path.exists(path):
        os.startfile(path)
    else:
        st.error(f"ファイルが見つかりません。\n{path}")

FASTAPI = "http://localhost:8000"

def get_pdf_token(filename):

    r = requests.post(
        f"{FASTAPI}/create_token",
        json={
            "filename": filename,
            "user": st.session_state["username"]
        },
        timeout=5
    )

    r.raise_for_status()

    return r.json()["token"]
def show_pdf(pdf_path, page):

    relative_filename = os.path.relpath(pdf_path, BASE_PDF)
    relative_filename = relative_filename.replace("\\", "/")

    token = get_pdf_token(relative_filename)

    pdf_url = (
        f"{FASTAPI}/view/{relative_filename}"
        f"?token={token}"
    )

    viewer_url = (
        f"{FASTAPI}/pdfjs/web/viewer.html"
        f"?file={quote(pdf_url, safe='')}"
        f"#page={page+1}"
    )

    # デバッグ用
    print(viewer_url)

    components.html(
        f"""
        <iframe
            src="{viewer_url}"
            width="100%"
            height="900"
            style="border:none;">
        </iframe>
        """,
        height=900,
    )

def bold(text):
    return f"<b>{text}</b>"
# =====================================================
# PowerPoint検索結果
# =====================================================
def score_star(score):

    if score >= 0.95:
        return "★★★★★"
    elif score >= 0.90:
        return "★★★★☆"
    elif score >= 0.85:
        return "★★★☆☆"
    elif score >= 0.80:
        return "★★☆☆☆"
    else:
        return "★☆☆☆☆"

# =====================================================
# UI
# =====================================================
st.set_page_config(
    page_title="Paper RAG",
    layout="wide"
)
# =====================================
# ジャンル選択
# =====================================
genre= st.sidebar.selectbox(
    "ジャンル",
    list(GENRE_SETTINGS.keys())
)
setting = GENRE_SETTINGS[genre]

DB_PATH = os.path.join(
    BASE_DB,
    setting["db"]
)

PDF_DIR = os.path.join(
    BASE_PDF,
    setting["pdf"]
)

TOP_K = setting["top_k"]

SYSTEM_PROMPT = setting["prompt"]

os.makedirs(DB_PATH, exist_ok=True)
os.makedirs(PDF_DIR, exist_ok=True)

st.title(f"📚 文書RAG（{genre}）")
if "uploader_key" not in st.session_state:
    st.session_state.uploader_key = 0
# =====================================================
# ChromaDB
# =====================================================
client = chromadb.PersistentClient(
    path=DB_PATH
)

collection = client.get_or_create_collection(
    name="documents"
)
# =====================================================
# question_log.DB
# =====================================================
question_conn = get_question_conn()
question_cur = question_conn.cursor()

# -----------------------------------------------------
# Upload
# -----------------------------------------------------
st.header("File 登録")

files = st.file_uploader(
    "アップロード",
    type=setting["file_types"],
    accept_multiple_files=True,
    key=f"pdf_uploader_{st.session_state.uploader_key}"
)

if st.button("登録開始"):

    if files:

        total = 0

        progress = st.progress(0)

        for i, file in enumerate(files):
            try:

                ext = os.path.splitext(file.name)[1].lower().replace(".", "")

                # ジャンルに対応していない拡張子は登録しない
                if ext not in setting["file_types"]:
                    st.warning(f"{file.name} はこのジャンルでは登録できません。")
                    continue

                # ファイル種別ごとの登録
                if ext == "pdf":
                    count = add_pdf(file)
                
                elif ext == "docx":
                    with tempfile.TemporaryDirectory() as tmpdir:

                        # Word保存
                        docx_path = os.path.join(tmpdir, os.path.basename(file.name))
                        with open(docx_path, "wb") as f:
                            f.write(file.getbuffer())

                        try:
                            if platform.system() == "Windows":
                                pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
                                convert(docx_path, pdf_path)

                            else:
                                subprocess.run(
                                    [
                                        "soffice",
                                        "--headless",
                                        "--convert-to", "pdf",
                                        "--outdir", tmpdir,
                                        docx_path
                                    ],
                                    check=True,
                                    stdout=subprocess.PIPE,
                                    stderr=subprocess.PIPE,
                                    text=True,
                                )

                                pdf_path = os.path.splitext(docx_path)[0] + ".pdf"

                            if not os.path.exists(pdf_path):
                                st.error(f"{file.name} のPDF変換に失敗しました。")
                                continue

                        except Exception as e:
                            st.error(f"{file.name} のPDF変換中にエラーが発生しました。")
                            st.exception(e)
                            continue

                        # Streamlit UploadedFile風オブジェクトを作る
                        with open(pdf_path, "rb") as f:
                            pdf_data = BytesIO(f.read())

                        pdf_data.name = os.path.splitext(file.name)[0] + ".pdf"

                        # 既存PDFルーチンへ
                        count = add_pdf(pdf_data) or 0
                
                elif ext == "pptx":
                    count = add_ppt(file)

                else:
                    st.warning(f"{file.name} は未対応のファイル形式です。")
                    continue

                total += count

            except Exception as e:
                st.error(f"{file.name} の処理中にエラーが発生。")
                st.exception(e)

            finally:
                # 成功・失敗・continueに関係なく必ず実行される
                progress.progress((i + 1) / len(files))
        
        show_upload_dialog(
            f"{len(files)}ファイル登録完了 "
            f"({total}チャンク)"
        )

# -----------------------------------------------------
# Ask
# -----------------------------------------------------
st.header("質問")
question = st.text_area("質問内容")
exclude_text = st.text_input(
    "🚫 検索対象から除外するファイル名（カンマ区切り）",
    placeholder="例：細胞接着, 古い資料, test.pdf"
)

# =========================
# 回答生成
# =========================
if "running" not in st.session_state:
    st.session_state.running = False

# ボタンが押されたら処理開始フラグを立てるだけ
if st.button("回答生成", disabled=st.session_state.running):

    # 前回のPDF表示を消す
    st.session_state.pop("selected_pdf", None)
    st.session_state.pop("selected_page", None)
    st.session_state.pop("selected_chunk_text", None)

    st.session_state.running = True
    st.rerun()

# =========================
# 回答生成本体
# =========================
if st.session_state.running:

    try:

        if question:
            start_time = time.perf_counter()

            with st.spinner("検索中..."):

                exclude_files = [
                    x.strip()
                    for x in re.split(r"[、,]", exclude_text)
                    if x.strip()
                ]

                answer, source_meta, source_doc, distances = answer_question(
                    question,
                    exclude_files
                )

            elapsed = time.perf_counter() - start_time

            # logDBに記録
            try:
                question_cur.execute(
                    """
                    INSERT INTO question_log
                    (datetime, username, genre, question, elapsed)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        st.session_state["username"],
                        genre,
                        question,
                        elapsed
                    )
                )
                question_conn.commit()

            except Exception as e:
                print(f"question_log登録失敗: {e}")

            st.session_state.answer = answer
            st.session_state.source_meta = source_meta
            st.session_state.source_doc = source_doc
            st.session_state.source_distance = distances

    finally:
        # 必ず解除
        st.session_state.running = False
        st.rerun()

# =========================
# 回答表示
# =========================
if "answer" in st.session_state:

    st.markdown("### 回答")
    st.write(st.session_state.answer)

    st_copy_to_clipboard(
        st.session_state.answer,
        "📋 回答をコピー"
    )

    st.markdown("---")

    # =====================================================
    # 選択されたPDFを最初に表示
    # =====================================================

    if (
        st.session_state.get("selected_pdf")
        and os.path.exists(st.session_state.selected_pdf)
    ):

        st.markdown("## 📖 PDFビューア（ジャンプ）")

        show_pdf(
            st.session_state["selected_pdf"],
            st.session_state.get("selected_page", 0)
        )

        st.markdown("---")

    # =====================================================
    # 参考文献
    # =====================================================

    st.markdown("## 📚 回答に用いられた資料")

    source_meta = st.session_state.source_meta
    source_doc = st.session_state.source_doc

    shown_files = []
    file_chunk_map = {}

    # =====================================================
    # メタデータ整理
    # =====================================================

    for meta in source_meta:

        file_name = meta["file"]
        page = meta.get("page", 0)
        chunk_id = meta.get("chunk", 0)

        if file_name not in shown_files:
            shown_files.append(file_name)

        if file_name not in file_chunk_map:
            file_chunk_map[file_name] = set()

        file_chunk_map[file_name].add(chunk_id)

    # =====================================================
    # PDF + PowerPoint チャンク表示
    # =====================================================

    for idx, meta in enumerate(source_meta):

        file_name = meta["file"]
        page = meta.get("page", 0)
        chunk_id = meta.get("chunk", 0)

        ext = os.path.splitext(file_name)[1].lower()

        # =================================================
        # PowerPoint
        # =================================================

        if ext == ".pptx":

            preview = source_doc[idx].strip()

            if len(preview) > 300:
                preview = preview[:300] + "\n..."

            score_text = ""

            if "source_distance" in st.session_state:
                score = 1 - st.session_state.source_distance[idx]
                score_text = f"{score_star(score)}　{score:.2f}"

            st.markdown(f"""
            {score_text}

            📄 **{file_name}**

            **Slide {page + 1}**

            {preview}

            """)

            st.divider()

        # =================================================
        # PDF
        # =================================================

        else:

            col1, col2 = st.columns([3, 1])

            with col1:

                st.markdown(
                    bold(source_doc[idx][:400]),
                    unsafe_allow_html=True
                )

                if st.button(
                    f"📄 {file_name} / chunk {chunk_id} / page {page + 1}",
                    key=f"btn_{file_name}_{chunk_id}_{page}"
                ):

                    st.session_state["selected_pdf"] = os.path.join(
                        PDF_DIR,
                        file_name
                    )

                    st.session_state["selected_page"] = page

                    st.session_state["selected_chunk_text"] = (
                        source_doc[idx]
                    )

                    st.rerun()

            with col2:

                st.write(f"p.{page + 1}")

    # =====================================================
    # ダウンロード
    # =====================================================

    st.caption(
        f"クリックでダウンロード　参考文献数: {len(shown_files)}"
    )

    for file_name in shown_files:

        pdf_path = os.path.join(PDF_DIR, file_name)

        if os.path.exists(pdf_path):

            with open(pdf_path, "rb") as f:
                pdf_bytes = f.read()

            st.download_button(
                f"📄 {file_name}",
                pdf_bytes,
                file_name=file_name,
                mime="application/pdf",
                key=f"dl_{file_name}"
            )
# =====================================
# PDF管理
# =====================================
st.sidebar.header("文書管理")

#-------------------------------
# 利用履歴
#-------------------------------
if st.sidebar.button("📊 利用履歴"):
    st.switch_page("pages/question_log.py")

st.sidebar.divider()

count = collection.count()

st.sidebar.metric(
    "登録チャンク数",
    count
)

# チャンクが1件以上ある場合のみ表示
if count > 0:

    all_data = collection.get()

    files = sorted(
        list(
            set(
                m["file"]
                for m in all_data["metadatas"]
                if "file" in m
            )
        )
    )

    if len(files) > 0:

        # ==========================
        # 文書絞り込み
        # ==========================
        filter_text = st.sidebar.text_input(
            "文書を絞り込み",
            placeholder="ファイル名を入力"
        )

        # ==========================
        # 入力文字を含むファイルだけ抽出
        # ==========================
        if filter_text:
            filtered_files = [
                f for f in files
                if filter_text.lower() in f.lower()
            ]
        else:
            filtered_files = files

        # ==========================
        # 削除対象文書
        # ==========================
        if len(filtered_files) > 0:

            target = st.sidebar.selectbox(
                "削除対象文書",
                filtered_files
            )

            if st.sidebar.button(
                "文書削除",
                type="primary"
            ):

                data = collection.get(
                    where={
                        "file": target
                    }
                )

                if len(data["ids"]) > 0:

                    # ==========================
                    # 表示中のPDFなら先にビューアを閉じる
                    # ==========================
                    current_pdf = st.session_state.get("selected_pdf")

                    if current_pdf and os.path.basename(current_pdf) == target:
                        st.session_state.pop("selected_pdf", None)
                        st.session_state.pop("selected_page", None)
                        st.session_state.pop("selected_chunk_text", None)

                    # ==========================
                    # ChromaDBから削除
                    # ==========================
                    collection.delete(ids=data["ids"])

                    # ==========================
                    # 保存されているPDFを削除
                    # ==========================
                    pdf_path = os.path.join(PDF_DIR, target)

                    if os.path.exists(pdf_path):
                        try:
                            os.remove(pdf_path)
                        except Exception as e:
                            st.sidebar.warning(
                                f"文書削除失敗: {e}"
                            )

                    st.sidebar.success(
                        f"{target} をDBと文書保存フォルダから削除しました"
                    )

                    st.rerun()

        else:

            st.sidebar.info(
                f"「{filter_text}」を含む文書はありません"
            )

else:

    st.sidebar.info(
        "登録文書はありません"
    )

